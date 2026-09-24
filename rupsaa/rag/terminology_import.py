"""Bulk CSV/XLSX import (and templates) for structured terminology.

Two-step, stateless workflow used by the owner UI (web/terminology.js) via
api/owner_routes.py:

  1. preview(file)  → parse + validate every row, classify duplicates, change nothing
  2. commit(file, update_rows) → re-parse + re-validate the SAME file, create the
     valid new rows, and UPDATE only the existing-term rows the owner explicitly
     chose (default for every duplicate is SKIP)

Re-validating on commit (instead of trusting a cached preview) means the
result can never drift from what the rules allow, and nothing is held
server-side between the two requests.

Records are written through TerminologyStore.create/update — the exact same
path as a manually created term — so imported terms are editable/deletable in
the Knowledge Manager, participate in routing, and are live on the next chat
message (no reindex, no retraining).

Safety: CSV is decoded as text only; XLSX is opened with openpyxl in
read-only, values-only mode (formulas are never evaluated — a formula cell
yields its cached value or nothing, and a text value starting with "=" is
flagged). Macro-enabled formats are rejected by extension. Size and row
limits are enforced before/while parsing.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import asdict, dataclass, field

from rupsaa.rag.terminology import (
    TERM_CATEGORIES,
    TERM_LANGUAGES,
    TerminologyError,
    TerminologyStore,
    TermRecord,
    normalize,
)

COLUMNS = [
    "term", "aliases", "category", "definition", "details", "answer_guidance",
    "languages", "example_queries", "tags", "enabled",
]
REQUIRED_COLUMNS = ["term", "definition"]
LIST_COLUMNS = ["aliases", "languages", "example_queries", "tags"]
SUPPORTED_EXTENSIONS = (".csv", ".xlsx")
MAX_FILE_BYTES = 2 * 1024 * 1024  # 2 MB
MAX_ROWS = 1000  # data rows (blank rows don't count)
MAX_CELL_CHARS = 6000
SHEET_NAME = "Terminology"

_TRUE = {"true", "yes", "y", "1"}
_FALSE = {"false", "no", "n", "0"}
_LANGUAGE_ALIASES = {
    "en": "en", "english": "en", "eng": "en",
    "bn": "bn", "bengali": "bn", "bangla": "bn", "বাংলা": "bn",
    "banglish": "banglish", "romanized bengali": "banglish", "romanised bengali": "banglish",
}

EXAMPLE_ROW = {
    "term": "Strip / Stripping",
    "aliases": "strip|stripping|strip ki|strip mane ki|স্ট্রিপ",
    "category": "Adult terminology",
    "definition": "Removing clothing, sometimes gradually, including in seductive, performance or sexual contexts depending on context.",
    "details": "A strip or stripping performance can involve gradually removing clothing. Meaning depends on context.",
    "answer_guidance": "Explain directly in the user's language/register. Start concise for simple definition questions and provide more detail when requested.",
    "languages": "en|banglish|bn",
    "example_queries": "strip mane ki|strip ki|স্ট্রিপ মানে কী|what does stripping mean",
    "tags": "adult|terminology|definition",
    "enabled": "true",
}

INSTRUCTIONS = [
    ("Rupsaa terminology import", ""),
    ("", ""),
    ("Main sheet", f'Fill the "{SHEET_NAME}" sheet: one term per row, first row = column headers (keep them as they are).'),
    ("Required columns", "term, definition"),
    ("Optional columns", "aliases, category, details, answer_guidance, languages, example_queries, tags, enabled"),
    ("Do NOT add", "id, created_at, updated_at — Rupsaa generates and manages these."),
    ("", ""),
    ("Multi-value columns", "aliases, languages, example_queries, tags: separate values with a pipe |"),
    ("Example", "strip|stripping|strip mane ki|স্ট্রিপ"),
    ("aliases", "Other ways people write the term — include Banglish and Bengali forms. These are what make lookups work."),
    ("category", "One of: " + ", ".join(TERM_CATEGORIES) + ' (friendly forms like "Adult terminology" are accepted). Blank = general.'),
    ("languages", "Any of: en, bn, banglish (also accepted: english, bengali, bangla). Blank = all three."),
    ("enabled", "true / false (also yes/no, 1/0). Blank = true. Disabled terms are stored but never used in chat."),
    ("", ""),
    ("Import behaviour", "Uploading shows a PREVIEW first — nothing is saved until you click \"Import Valid Records\"."),
    ("Invalid rows", "Rows with errors are listed with their row number and skipped; they never block the valid rows."),
    ("Blank rows", "Ignored."),
    ("Duplicates in the file", "If two rows share a term/alias (ignoring case, spaces, punctuation), only the first is imported."),
    ("Existing terms", "A row matching an existing term defaults to SKIP. Choose UPDATE in the preview to overwrite it: "
                       "the record keeps its id and created date; non-blank columns replace the stored values, blank columns keep them."),
    ("After import", "Terms are live in chat on the next message — no reindex, no retraining — and editable in the Terminology tab."),
    ("", ""),
    ("Limits", f".csv or .xlsx only, max {MAX_FILE_BYTES // (1024 * 1024)} MB, max {MAX_ROWS} rows. Formulas are not evaluated; macros are not supported."),
]


class ImportFileError(ValueError):
    """The whole file is unusable (bad type, too big, missing columns, …)."""


@dataclass
class ImportRow:
    row: int  # spreadsheet row number (header = row 1)
    status: str  # valid | warning | invalid | duplicate_in_file | existing_match
    data: dict = field(default_factory=dict)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    existing_id: str | None = None
    existing_term: str | None = None
    duplicate_of_row: int | None = None


@dataclass
class ImportPreview:
    file_type: str
    total_rows: int
    rows: list[ImportRow]

    @property
    def counts(self) -> dict:
        c = {"total": self.total_rows, "valid": 0, "warnings": 0, "invalid": 0, "duplicates": 0,
             "duplicates_in_file": 0, "existing_matches": 0}
        for r in self.rows:
            if r.status in ("valid", "warning"):
                c["valid"] += 1
            if r.warnings:
                c["warnings"] += 1
            if r.status == "invalid":
                c["invalid"] += 1
            if r.status == "duplicate_in_file":
                c["duplicates_in_file"] += 1
            if r.status == "existing_match":
                c["existing_matches"] += 1
        c["duplicates"] = c["duplicates_in_file"] + c["existing_matches"]
        return c

    def to_dict(self) -> dict:
        return {"file_type": self.file_type, "counts": self.counts, "rows": [asdict(r) for r in self.rows]}


# --- file reading --------------------------------------------------------------------

def _check_file(filename: str, content: bytes) -> str:
    name = (filename or "").strip().lower()
    ext = next((e for e in SUPPORTED_EXTENSIONS if name.endswith(e)), None)
    if ext is None:
        raise ImportFileError("Unsupported file type — upload a .csv or .xlsx file.")
    if not content:
        raise ImportFileError("The file is empty.")
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError(f"File is too large (max {MAX_FILE_BYTES // (1024 * 1024)} MB).")
    if ext == ".xlsx" and not zipfile.is_zipfile(io.BytesIO(content)):
        raise ImportFileError("This is not a valid .xlsx file (it may be an old .xls or a renamed file).")
    return ext


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value)


def _read_csv(content: bytes) -> list[list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ImportFileError("CSV must be saved as UTF-8 (Excel: \"CSV UTF-8 (Comma delimited)\").")
    try:
        return [row for row in csv.reader(io.StringIO(text, newline=""))]
    except csv.Error as e:
        raise ImportFileError(f"Could not parse CSV: {e}")


def _read_xlsx(content: bytes) -> list[list[str]]:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    except Exception:
        raise ImportFileError("Could not open the .xlsx file — is it a valid Excel workbook?")
    try:
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[0]
        rows = []
        for i, values in enumerate(ws.iter_rows(values_only=True)):
            if i > MAX_ROWS + 1 + 5000:  # hard stop on absurd sheets (blank rows included)
                break
            rows.append([_cell(v) for v in values])
        return rows
    finally:
        wb.close()


# --- row normalization ----------------------------------------------------------------

def _split_pipe(value: str) -> list[str]:
    out, seen = [], set()
    for part in value.split("|"):
        part = part.strip()
        if part and normalize(part) not in seen:
            seen.add(normalize(part))
            out.append(part)
    return out


def _category(value: str, warnings: list[str]) -> str:
    if not value:
        return "general"
    key = re.sub(r"[\s\-/]+", "_", value.strip().lower())
    if key in TERM_CATEGORIES:
        return key
    for cat in TERM_CATEGORIES:  # "creator platform terms" → creator_platform, "adult" → adult_terminology
        if len(key) >= 4 and (key.startswith(cat) or cat.startswith(key)):
            return cat
    warnings.append(f"Unknown category {value!r} — using 'general' (allowed: {', '.join(TERM_CATEGORIES)})")
    return "general"


def _languages(values: list[str], warnings: list[str]) -> list[str]:
    out = []
    for v in values:
        lang = _LANGUAGE_ALIASES.get(v.strip().lower())
        if lang is None:
            warnings.append(f"Unknown language {v!r} ignored (allowed: {', '.join(TERM_LANGUAGES)})")
        elif lang not in out:
            out.append(lang)
    return out or list(TERM_LANGUAGES)


def _parse_row(row_number: int, raw: dict[str, str]) -> ImportRow:
    errors: list[str] = []
    warnings: list[str] = []
    cells = {k: (raw.get(k) or "").strip() for k in COLUMNS}
    for key, value in cells.items():
        if len(value) > MAX_CELL_CHARS:
            errors.append(f"{key} is longer than {MAX_CELL_CHARS} characters")
        if value.startswith("="):
            warnings.append(f"{key} starts with '=' — treated as plain text (formulas are not evaluated)")
    for key in REQUIRED_COLUMNS:
        if not cells[key]:
            errors.append(f"Missing required field: {key}")

    enabled_raw = cells["enabled"].lower()
    if not enabled_raw:
        enabled = True
    elif enabled_raw in _TRUE:
        enabled = True
    elif enabled_raw in _FALSE:
        enabled = False
    else:
        enabled = True
        errors.append(f"Invalid enabled value {cells['enabled']!r} (use true or false)")

    data = {
        "term": cells["term"],
        "definition": cells["definition"],
        "category": _category(cells["category"], warnings),
        "aliases": _split_pipe(cells["aliases"]),
        "details": cells["details"],
        "answer_guidance": cells["answer_guidance"],
        "languages": _languages(_split_pipe(cells["languages"]), warnings),
        "example_queries": _split_pipe(cells["example_queries"]),
        "tags": _split_pipe(cells["tags"]),
        "enabled": enabled,
    }
    # Which columns actually had content (UPDATE only overwrites these).
    data["_provided"] = [k for k in COLUMNS if cells[k]]
    if data["term"] and len(data["term"]) > 120:
        errors.append("term is longer than 120 characters")
    if data["definition"] and len(data["definition"]) > 2000:
        errors.append("definition is longer than 2000 characters")
    status = "invalid" if errors else ("warning" if warnings else "valid")
    return ImportRow(row=row_number, status=status, data=data, errors=errors, warnings=warnings)


def _row_keys(data: dict) -> set[str]:
    rec = TermRecord(id="term-x", term=data["term"], definition=data["definition"], aliases=data["aliases"])
    return rec.keys()


# --- preview / commit ---------------------------------------------------------------------

def preview(filename: str, content: bytes, store: TerminologyStore) -> ImportPreview:
    ext = _check_file(filename, content)
    table = _read_csv(content) if ext == ".csv" else _read_xlsx(content)
    if not table or not any(c.strip() for c in table[0]):
        raise ImportFileError("The file has no header row.")

    header = [normalize(h).replace(" ", "_") for h in table[0]]
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        raise ImportFileError(f"Missing required column(s): {', '.join(missing)}. "
                              f"Expected headers: {', '.join(COLUMNS)}")
    unknown = [table[0][i] for i, h in enumerate(header) if h and h not in COLUMNS]

    rows: list[ImportRow] = []
    data_rows = 0
    for idx, values in enumerate(table[1:], start=2):
        if not any((v or "").strip() for v in values):
            continue  # completely blank row
        data_rows += 1
        if data_rows > MAX_ROWS:
            raise ImportFileError(f"Too many rows — the limit is {MAX_ROWS} terms per file.")
        raw = {header[i]: values[i] for i in range(min(len(header), len(values))) if header[i] in COLUMNS}
        rows.append(_parse_row(idx, raw))
    if unknown and rows:
        rows[0].warnings.append(f"Ignored unknown column(s): {', '.join(unknown)}")
        if rows[0].status == "valid":
            rows[0].status = "warning"

    # Duplicates: inside the file, then against existing terminology.
    existing = store.list()
    seen: dict[str, int] = {}
    for r in rows:
        if r.status == "invalid":
            continue
        keys = _row_keys(r.data)
        clash = next((seen[k] for k in keys if k in seen), None)
        if clash is not None:
            r.status = "duplicate_in_file"
            r.duplicate_of_row = clash
            r.errors.append(f"Duplicate of row {clash} in this file (same term or alias) — not imported")
            continue
        for k in keys:
            seen[k] = r.row
        matches = [e for e in existing if keys & e.keys()]
        if len(matches) == 1:
            r.status = "existing_match"
            r.existing_id, r.existing_term = matches[0].id, matches[0].term
        elif len(matches) > 1:
            r.status = "invalid"
            r.errors.append("Term/aliases match several existing terms ("
                            + ", ".join(m.term for m in matches) + ") — edit those terms manually")
    return ImportPreview(file_type=ext.lstrip("."), total_rows=data_rows, rows=rows)


@dataclass
class ImportResult:
    created: list[dict] = field(default_factory=list)
    updated: list[dict] = field(default_factory=list)
    skipped: list[dict] = field(default_factory=list)
    failed: list[dict] = field(default_factory=list)

    def to_dict(self) -> dict:
        d = asdict(self)
        d["counts"] = {k: len(v) for k, v in d.items()}
        return d


def commit(filename: str, content: bytes, store: TerminologyStore, update_rows: set[int] | None = None) -> ImportResult:
    """Import valid rows. Existing matches are SKIPPED unless their row number
    is in `update_rows` (explicit owner choice)."""
    update_rows = set(update_rows or ())
    result = ImportResult()
    for r in preview(filename, content, store).rows:
        data = {k: v for k, v in r.data.items() if k != "_provided"}
        if r.status in ("invalid", "duplicate_in_file"):
            result.skipped.append({"row": r.row, "term": r.data.get("term"), "reason": "; ".join(r.errors)})
            continue
        try:
            if r.status == "existing_match":
                if r.row not in update_rows:
                    result.skipped.append({"row": r.row, "term": data["term"],
                                           "reason": f"already exists as {r.existing_id} (SKIP)"})
                    continue
                provided = set(r.data["_provided"])
                changes = {k: data[k] for k in ("term", "definition", "category", "aliases", "details",
                                                "answer_guidance", "languages", "example_queries", "tags", "enabled")
                           if k in provided}
                rec = store.update(r.existing_id, changes)
                result.updated.append({"row": r.row, "id": rec.id, "term": rec.term})
            else:
                rec = store.create(data)
                result.created.append({"row": r.row, "id": rec.id, "term": rec.term})
        except TerminologyError as e:
            result.failed.append({"row": r.row, "term": data.get("term"), "error": str(e)})
    return result


# --- templates -------------------------------------------------------------------------------

def template_csv() -> bytes:
    buf = io.StringIO(newline="")
    writer = csv.DictWriter(buf, fieldnames=COLUMNS, lineterminator="\r\n")
    writer.writeheader()
    writer.writerow(EXAMPLE_ROW)
    # UTF-8 BOM so Excel shows Bengali correctly when the CSV is double-clicked.
    return ("﻿" + buf.getvalue()).encode("utf-8")


def template_xlsx() -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(COLUMNS)
    ws.append([EXAMPLE_ROW[c] for c in COLUMNS])
    header_fill = PatternFill("solid", fgColor="F3E1EC")
    for col_idx, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="B00020" if name in REQUIRED_COLUMNS else "000000")
        cell.fill = header_fill
        ws.column_dimensions[cell.column_letter].width = {"definition": 60, "details": 50, "answer_guidance": 50,
                                                          "aliases": 40, "example_queries": 45}.get(name, 18)
        # Every cell is text: stops Excel turning "true" or long numbers into other types.
        for row in range(2, 202):
            ws.cell(row=row, column=col_idx).number_format = "@"
    for cell in ws[2]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"

    info = wb.create_sheet("Instructions")
    for a, b in INSTRUCTIONS:
        info.append([a, b])
    info["A1"].font = Font(bold=True, size=14)
    for row in info.iter_rows(min_row=3):
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
