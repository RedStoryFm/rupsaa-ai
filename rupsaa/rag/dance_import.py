"""Bulk import of owner dance-style records (CSV / XLSX / JSON) into the DanceStore.

Same preview → commit flow as terminology import (rupsaa/rag/terminology_import.py):
preview validates every row and classifies it (valid / warning / invalid /
duplicate_in_file / existing_match); commit creates valid rows and SKIPS existing
dances unless the owner explicitly asks to update those rows.

Only what the owner supplies is stored — nothing is generated or filled in.
Header names are matched loosely ("Dance Style", "Region", "Also known as" …).
"""

from __future__ import annotations

import csv
import io
import json
from dataclasses import asdict, dataclass, field

from rupsaa.rag.dance import DANCE_LANGUAGES, FUTURE_FIELDS, DanceError, DanceRecord, DanceStore
from rupsaa.rag.terminology import normalize
from rupsaa.rag.terminology_import import (
    MAX_CELL_CHARS,
    MAX_FILE_BYTES,
    ImportFileError,
    _cell,
    _read_csv,
    _read_xlsx,
    _split_pipe,
)

COLUMNS = ["name", "aliases", "origin", "category", "description", "key_movements", "answer_guidance",
           "languages", "tags", "enabled", *FUTURE_FIELDS, "source"]
REQUIRED_COLUMNS = ["name", "description"]
SUPPORTED_EXTENSIONS = (".csv", ".xlsx", ".json")
MAX_ROWS = 1000
SHEET_NAME = "Dance"
HEADER_SYNONYMS = {
    "name": ["name", "dance", "dance_name", "dance_style", "style", "dance_form", "title"],
    "aliases": ["aliases", "alias", "also_known_as", "aka", "other_names", "alternate_names"],
    "origin": ["origin", "region", "origin_region", "country", "country_of_origin", "place_of_origin", "originated"],
    "category": ["category", "type", "genre", "family"],
    "description": ["description", "about", "summary", "details", "overview"],
    "key_movements": ["key_movements", "movements", "key_moves", "moves", "signature_moves"],
    "answer_guidance": ["answer_guidance", "guidance"],
    "languages": ["languages", "language"],
    "tags": ["tags", "tag"],
    "enabled": ["enabled", "active"],
    "source": ["source", "notes"],
    **{f: [f] for f in FUTURE_FIELDS},
}
_HEADER_MAP = {normalize(s).replace(" ", "_"): canon for canon, syns in HEADER_SYNONYMS.items() for s in syns}
_TRUE, _FALSE = {"true", "yes", "y", "1"}, {"false", "no", "n", "0"}
_LANGUAGE_ALIASES = {"en": "en", "english": "en", "bn": "bn", "bengali": "bn", "bangla": "bn", "বাংলা": "bn",
                     "banglish": "banglish", "romanized bengali": "banglish"}
PLACEHOLDER_ROW = {"name": "Example Dance (replace me)", "aliases": "example dance|sample dance",
                   "origin": "(where the owner says it comes from)", "category": "(owner's category)",
                   "description": "(owner's description — only what you know to be true)",
                   "key_movements": "(optional)", "enabled": "false"}


@dataclass
class DanceImportRow:
    row: int
    status: str  # valid | warning | invalid | duplicate_in_file | existing_match
    data: dict = field(default_factory=dict)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    existing_id: str | None = None
    existing_name: str | None = None
    duplicate_of_row: int | None = None


@dataclass
class DanceImportPreview:
    file_type: str
    total_rows: int
    rows: list[DanceImportRow]

    @property
    def counts(self) -> dict:
        c = {"total": self.total_rows, "valid": 0, "warnings": 0, "invalid": 0, "duplicates_in_file": 0, "existing_matches": 0,
             "missing_origin": 0}
        for r in self.rows:
            c["valid"] += r.status in ("valid", "warning")
            c["warnings"] += bool(r.warnings)
            c["invalid"] += r.status == "invalid"
            c["duplicates_in_file"] += r.status == "duplicate_in_file"
            c["existing_matches"] += r.status == "existing_match"
            c["missing_origin"] += not r.data.get("origin")
        return c

    def to_dict(self) -> dict:
        return {"file_type": self.file_type, "counts": self.counts, "rows": [asdict(r) for r in self.rows]}


def _check_file(filename: str, content: bytes) -> str:
    name = (filename or "").strip().lower()
    ext = next((e for e in SUPPORTED_EXTENSIONS if name.endswith(e)), None)
    if ext is None:
        raise ImportFileError("Unsupported file type — upload a .csv, .xlsx or .json file.")
    if not content:
        raise ImportFileError("The file is empty.")
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError(f"File is too large (max {MAX_FILE_BYTES // (1024 * 1024)} MB).")
    return ext


def _read_json(content: bytes) -> list[dict]:
    try:
        data = json.loads(content.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as e:
        raise ImportFileError(f"Could not parse JSON: {e}")
    if isinstance(data, dict):
        data = data.get("dances") or data.get("records") or data.get("items")
    if not isinstance(data, list) or not all(isinstance(x, dict) for x in data):
        raise ImportFileError('JSON must be a list of objects (or {"dances": [...]}).')
    return data


def _table_to_dicts(table: list[list[str]]) -> tuple[list[dict], list[str]]:
    if not table or not any(c.strip() for c in table[0]):
        raise ImportFileError("The file has no header row.")
    header = [_HEADER_MAP.get(normalize(h).replace(" ", "_")) for h in table[0]]
    unknown = [table[0][i] for i, h in enumerate(header) if h is None and table[0][i].strip()]
    rows = []
    for values in table[1:]:
        rows.append({header[i]: values[i] for i in range(min(len(header), len(values))) if header[i]})
    return rows, unknown


def _as_text(value) -> str:
    if isinstance(value, list):
        return " | ".join(str(v) for v in value)
    return _cell(value)


def _parse(row_number: int, raw: dict) -> DanceImportRow:
    errors: list[str] = []
    warnings: list[str] = []
    cells = {k: _as_text(raw.get(k)).strip() for k in COLUMNS}
    for key, value in cells.items():
        if len(value) > MAX_CELL_CHARS:
            errors.append(f"{key} is longer than {MAX_CELL_CHARS} characters")
        if value.startswith("="):
            warnings.append(f"{key} starts with '=' — treated as plain text")
    for key in REQUIRED_COLUMNS:
        if not cells[key]:
            errors.append(f"Missing required field: {key}")
    if not cells["origin"]:
        warnings.append("No origin given — stored without one (nothing is filled in automatically)")
    enabled_raw = cells["enabled"].lower()
    enabled = True if not enabled_raw or enabled_raw in _TRUE else (False if enabled_raw in _FALSE else None)
    if enabled is None:
        errors.append(f"Invalid enabled value {cells['enabled']!r} (use true or false)")
        enabled = True
    langs = []
    for v in _split_pipe(cells["languages"].replace(",", "|")):
        lang = _LANGUAGE_ALIASES.get(v.lower())
        if lang is None:
            warnings.append(f"Unknown language {v!r} ignored (allowed: {', '.join(DANCE_LANGUAGES)})")
        elif lang not in langs:
            langs.append(lang)
    data = {k: cells[k] for k in COLUMNS if k not in ("aliases", "languages", "tags", "enabled")}
    data.update({"aliases": _split_pipe(cells["aliases"].replace(",", "|")), "tags": _split_pipe(cells["tags"].replace(",", "|")),
                 "languages": langs or list(DANCE_LANGUAGES), "enabled": enabled})
    data["_provided"] = [k for k in COLUMNS if cells[k]]
    if len(data["name"]) > 120:
        errors.append("name is longer than 120 characters")
    status = "invalid" if errors else ("warning" if warnings else "valid")
    return DanceImportRow(row=row_number, status=status, data=data, errors=errors, warnings=warnings)


def preview(filename: str, content: bytes, store: DanceStore) -> DanceImportPreview:
    ext = _check_file(filename, content)
    unknown: list[str] = []
    if ext == ".json":
        raw_rows = [{_HEADER_MAP.get(normalize(k).replace(" ", "_"), None): v for k, v in obj.items()} for obj in _read_json(content)]
        raw_rows = [{k: v for k, v in r.items() if k} for r in raw_rows]
        first_row = 1
    else:
        table = _read_csv(content) if ext == ".csv" else _read_xlsx(content)
        raw_rows, unknown = _table_to_dicts(table)
        first_row = 2
        if "name" not in {k for r in raw_rows for k in r} and raw_rows:
            raise ImportFileError("No 'name' column found (accepted: " + ", ".join(HEADER_SYNONYMS["name"]) + ").")
    rows: list[DanceImportRow] = []
    for idx, raw in enumerate(raw_rows, start=first_row):
        if not any(_as_text(v).strip() for v in raw.values()):
            continue
        if len(rows) >= MAX_ROWS:
            raise ImportFileError(f"Too many rows — the limit is {MAX_ROWS} dances per file.")
        rows.append(_parse(idx, raw))
    if unknown and rows:
        rows[0].warnings.append(f"Ignored unknown column(s): {', '.join(unknown)}")
        if rows[0].status == "valid":
            rows[0].status = "warning"

    existing = store.list()
    seen: dict[str, int] = {}
    for r in rows:
        if r.status == "invalid":
            continue
        keys = DanceRecord(id="dance-x", name=r.data["name"], description="x", aliases=r.data["aliases"]).keys()
        clash = next((seen[k] for k in keys if k in seen), None)
        if clash is not None:
            r.status, r.duplicate_of_row = "duplicate_in_file", clash
            r.errors.append(f"Duplicate of row {clash} in this file (same name or alias) — not imported")
            continue
        for k in keys:
            seen[k] = r.row
        matches = [e for e in existing if keys & e.keys()]
        if len(matches) == 1:
            r.status, r.existing_id, r.existing_name = "existing_match", matches[0].id, matches[0].name
        elif len(matches) > 1:
            r.status = "invalid"
            r.errors.append("Name/aliases match several existing dances (" + ", ".join(m.name for m in matches) + ")")
    return DanceImportPreview(file_type=ext.lstrip("."), total_rows=len(rows), rows=rows)


@dataclass
class DanceImportResult:
    created: list[dict] = field(default_factory=list)
    updated: list[dict] = field(default_factory=list)
    skipped: list[dict] = field(default_factory=list)
    failed: list[dict] = field(default_factory=list)

    def to_dict(self) -> dict:
        d = asdict(self)
        d["counts"] = {k: len(v) for k, v in d.items()}
        return d


def commit(filename: str, content: bytes, store: DanceStore, update_rows: set[int] | None = None) -> DanceImportResult:
    update_rows = set(update_rows or ())
    result = DanceImportResult()
    for r in preview(filename, content, store).rows:
        data = {k: v for k, v in r.data.items() if k != "_provided"}
        if r.status in ("invalid", "duplicate_in_file"):
            result.skipped.append({"row": r.row, "name": data.get("name"), "reason": "; ".join(r.errors)})
            continue
        try:
            if r.status == "existing_match":
                if r.row not in update_rows:
                    result.skipped.append({"row": r.row, "name": data["name"], "reason": f"already exists as {r.existing_id} (SKIP)"})
                    continue
                rec = store.update(r.existing_id, {k: data[k] for k in r.data["_provided"] if k in data})
                result.updated.append({"row": r.row, "id": rec.id, "name": rec.name})
            else:
                rec = store.create(data)
                result.created.append({"row": r.row, "id": rec.id, "name": rec.name})
        except DanceError as e:
            result.failed.append({"row": r.row, "name": data.get("name"), "error": str(e)})
    return result


def template_csv() -> bytes:
    buf = io.StringIO(newline="")
    writer = csv.DictWriter(buf, fieldnames=COLUMNS, lineterminator="\r\n")
    writer.writeheader()
    writer.writerow({k: PLACEHOLDER_ROW.get(k, "") for k in COLUMNS})
    return ("﻿" + buf.getvalue()).encode("utf-8")


def template_xlsx() -> bytes:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.append([PLACEHOLDER_ROW.get(k, "") for k in COLUMNS])
    notes = wb.create_sheet("How to fill")
    for line in ("Required: name, description. Everything else is optional.",
                 "Separate several aliases / tags / languages with |  (e.g. breakdance|b-boying).",
                 "Leave step fields empty unless you have verified steps — Rupsaa never invents them.",
                 "Rows whose name/alias already exists are skipped unless you choose UPDATE in the preview."):
        notes.append([line])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

