"""Terminology bulk import (CSV/XLSX), templates, duplicate SKIP/UPDATE.

Every test uses an isolated TerminologyStore in tmp_path (monkeypatched into
api.owner_routes), so the real knowledge/terminology/ is never touched. No
model or embedding model is loaded.
"""

from __future__ import annotations

import csv
import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

import api.owner_routes as owner_routes
from api.main import app
from rupsaa.rag import terminology_import as ti
from rupsaa.rag.context_builder import build_turn_knowledge
from rupsaa.rag.router import classify_message
from rupsaa.rag.terminology import TerminologyStore

HEADER = ti.COLUMNS
STRIP_ROW = dict(ti.EXAMPLE_ROW)
GFE_ROW = {
    "term": "GFE", "aliases": "girlfriend experience|জিএফই", "category": "adult_terminology",
    "definition": "Girlfriend experience: companionship that feels like a relationship.",
    "details": "", "answer_guidance": "", "languages": "en|banglish", "example_queries": "gfe mane ki",
    "tags": "adult", "enabled": "",
}


def csv_bytes(rows: list[dict], header: list[str] = HEADER) -> bytes:
    buf = io.StringIO(newline="")
    w = csv.DictWriter(buf, fieldnames=header, extrasaction="ignore")
    w.writeheader()
    for r in rows:
        w.writerow(r)
    return buf.getvalue().encode("utf-8")


def xlsx_bytes(rows: list[list], header: list[str] = HEADER, sheet: str = "Terminology") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(header)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def store(tmp_path, monkeypatch) -> TerminologyStore:
    s = TerminologyStore(tmp_path / "terms")
    monkeypatch.setattr(owner_routes, "_get_terminology", lambda: s)
    return s


@pytest.fixture
def client() -> TestClient:
    return TestClient(app)


def preview(client, name, content, **kw):
    return client.post("/owner/terminology/import/preview", files={"file": (name, content)}, **kw)


def commit(client, name, content, update_rows="", **kw):
    return client.post("/owner/terminology/import", files={"file": (name, content)},
                       data={"update_rows": update_rows}, **kw)


# --- templates -------------------------------------------------------------------------

def test_csv_template_download(client):
    r = client.get("/owner/terminology/template.csv")
    assert r.status_code == 200 and "terminology_template.csv" in r.headers["content-disposition"]
    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert list(rows[0].keys()) == HEADER and len(rows) == 1
    assert rows[0]["term"] == "Strip / Stripping" and "|" in rows[0]["aliases"] and "স্ট্রিপ" in rows[0]["aliases"]


def test_xlsx_template_download(client):
    r = client.get("/owner/terminology/template.xlsx")
    assert r.status_code == 200 and "terminology_template.xlsx" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    assert wb.sheetnames == ["Terminology", "Instructions"]
    rows = list(wb["Terminology"].iter_rows(values_only=True))
    assert list(rows[0]) == HEADER and rows[1][0] == "Strip / Stripping" and len([r for r in rows if any(r)]) == 2
    text = " ".join(str(c) for row in wb["Instructions"].iter_rows(values_only=True) for c in row if c)
    for needle in ("Required columns", "Optional columns", "pipe", "true / false", "SKIP", "UPDATE"):
        assert needle in text


def test_templates_round_trip_through_import(store, client):
    for name, url in (("t.csv", "/owner/terminology/template.csv"), ("t.xlsx", "/owner/terminology/template.xlsx")):
        body = preview(client, name, client.get(url).content).json()
        assert body["counts"]["valid"] == 1 and body["counts"]["invalid"] == 0


# --- preview -----------------------------------------------------------------------------

def test_valid_csv_preview_changes_nothing(store, client):
    r = preview(client, "terms.csv", csv_bytes([STRIP_ROW, GFE_ROW]))
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["file_type"] == "csv"
    assert body["counts"] == {"total": 2, "valid": 2, "warnings": 0, "invalid": 0, "duplicates": 0,
                              "duplicates_in_file": 0, "existing_matches": 0}
    row = body["rows"][0]
    assert row["row"] == 2 and row["data"]["category"] == "adult_terminology"  # "Adult terminology" normalized
    assert store.list() == []  # preview never writes


def test_valid_xlsx_preview(store, client):
    content = xlsx_bytes([[STRIP_ROW[c] for c in HEADER], [GFE_ROW[c] for c in HEADER]])
    body = preview(client, "terms.xlsx", content).json()
    assert body["file_type"] == "xlsx" and body["counts"]["valid"] == 2
    assert store.list() == []


def test_missing_required_columns(store, client):
    r = preview(client, "t.csv", csv_bytes([{"term": "x"}], header=["term", "aliases"]))
    assert r.status_code == 400 and "definition" in r.json()["detail"]


def test_missing_term_and_definition_reported_with_row_numbers(store, client):
    rows = [STRIP_ROW, {**GFE_ROW, "term": ""}, {**GFE_ROW, "term": "PPV", "aliases": "", "definition": "  "}]
    body = preview(client, "t.csv", csv_bytes(rows)).json()
    by_row = {r["row"]: r for r in body["rows"]}
    assert by_row[3]["status"] == "invalid" and "Missing required field: term" in by_row[3]["errors"]
    assert by_row[4]["status"] == "invalid" and "Missing required field: definition" in by_row[4]["errors"]
    assert by_row[2]["status"] == "valid" and body["counts"]["invalid"] == 2


def test_blank_rows_ignored_and_whitespace_trimmed(store, client):
    content = csv_bytes([{c: "" for c in HEADER}, {**GFE_ROW, "term": "  GFE  "}, {c: "   " for c in HEADER}])
    body = preview(client, "t.csv", content).json()
    assert body["counts"]["total"] == 1
    assert body["rows"][0]["row"] == 3 and body["rows"][0]["data"]["term"] == "GFE"


def test_pipe_list_parsing(store, client):
    row = {**GFE_ROW, "aliases": " a | b ||B| c ", "languages": "English| bangla |banglish",
           "tags": "x|y", "example_queries": "q1 | q2"}
    data = preview(client, "t.csv", csv_bytes([row])).json()["rows"][0]["data"]
    assert data["aliases"] == ["a", "b", "c"]  # trimmed, blanks dropped, case-insensitive dedupe
    assert data["languages"] == ["en", "bn", "banglish"]
    assert data["tags"] == ["x", "y"] and data["example_queries"] == ["q1", "q2"]


@pytest.mark.parametrize("value,expected,ok", [("true", True, True), ("FALSE", False, True), ("", True, True),
                                               ("yes", True, True), ("0", False, True), ("maybe", None, False)])
def test_enabled_parsing(store, client, value, expected, ok):
    row = preview(client, "t.csv", csv_bytes([{**GFE_ROW, "enabled": value}])).json()["rows"][0]
    if ok:
        assert row["status"] == "valid" and row["data"]["enabled"] is expected
    else:
        assert row["status"] == "invalid" and "Invalid enabled value" in row["errors"][0]


def test_xlsx_boolean_and_number_cells(store, client):
    content = xlsx_bytes([["PPV", "pay per view", "creator platform", "Pay-per-view content.", None, None,
                           None, None, None, False]])
    row = preview(client, "t.xlsx", content).json()["rows"][0]
    assert row["data"]["enabled"] is False and row["data"]["category"] == "creator_platform"


def test_unknown_category_and_language_are_warnings_not_errors(store, client):
    row = preview(client, "t.csv", csv_bytes([{**GFE_ROW, "category": "Spicy", "languages": "klingon|en"}])).json()["rows"][0]
    assert row["status"] == "warning" and row["data"]["category"] == "general" and row["data"]["languages"] == ["en"]
    assert len(row["warnings"]) == 2


# --- duplicates ------------------------------------------------------------------------------

def test_duplicate_within_upload(store, client):
    rows = [STRIP_ROW, {**GFE_ROW, "term": "  strip / STRIPPING "}, {**GFE_ROW, "term": "Lap dance", "aliases": "STRIP"}]
    body = preview(client, "t.csv", csv_bytes(rows)).json()
    assert [r["status"] for r in body["rows"]] == ["valid", "duplicate_in_file", "duplicate_in_file"]
    assert body["rows"][1]["duplicate_of_row"] == 2 and body["counts"]["duplicates_in_file"] == 2
    result = commit(client, "t.csv", csv_bytes(rows)).json()
    assert result["counts"]["created"] == 1 and len(store.list()) == 1


def test_duplicate_against_existing_defaults_to_skip(store, client):
    existing = store.create({"term": "strip", "definition": "Old definition.", "category": "adult_terminology"})
    body = preview(client, "t.csv", csv_bytes([STRIP_ROW, GFE_ROW])).json()
    assert body["rows"][0]["status"] == "existing_match" and body["rows"][0]["existing_id"] == existing.id
    result = commit(client, "t.csv", csv_bytes([STRIP_ROW, GFE_ROW])).json()  # no update_rows → SKIP
    assert result["counts"] == {"created": 1, "updated": 0, "skipped": 1, "failed": 0}
    assert store.get(existing.id).definition == "Old definition."
    assert len(store.list()) == 2  # no duplicate record for case/spacing differences


def test_explicit_update_preserves_id_and_created_at(store, client):
    existing = store.create({"term": "Strip", "definition": "Old definition.", "category": "adult_terminology",
                             "details": "keep me"})
    content = csv_bytes([{**STRIP_ROW, "details": ""}])
    result = commit(client, "t.csv", content, update_rows="2").json()
    assert result["counts"]["updated"] == 1 and result["updated"][0]["id"] == existing.id
    after = store.get(existing.id)
    assert after.id == existing.id and after.created_at == existing.created_at
    assert after.updated_at > existing.updated_at
    assert after.definition == STRIP_ROW["definition"] and after.term == "Strip / Stripping"
    assert after.details == "keep me"  # blank column keeps the stored value
    assert "স্ট্রিপ" in after.aliases and len(store.list()) == 1


def test_row_matching_several_existing_terms_is_invalid(store, client):
    store.create({"term": "Strip", "definition": "a"})
    store.create({"term": "GFE", "definition": "b"})
    body = preview(client, "t.csv", csv_bytes([{**GFE_ROW, "term": "Strip", "aliases": "gfe"}])).json()
    assert body["rows"][0]["status"] == "invalid" and "several existing terms" in body["rows"][0]["errors"][0]


# --- robustness / limits / security -------------------------------------------------------------

def test_malformed_rows_do_not_block_valid_rows(store, client):
    rows = [{**GFE_ROW, "definition": ""}, STRIP_ROW, {**GFE_ROW, "term": "PPV", "aliases": "", "enabled": "perhaps"}]
    result = commit(client, "t.csv", csv_bytes(rows)).json()
    assert result["counts"]["created"] == 1 and result["counts"]["skipped"] == 2
    assert [t.term for t in store.list()] == ["Strip / Stripping"]


@pytest.mark.parametrize("name,content", [("terms.xls", b"\xd0\xcf\x11\xe0junk"), ("terms.xlsm", b"PK\x03\x04"),
                                          ("terms.txt", b"term,definition\nx,y"), ("terms.json", b"{}")])
def test_unsupported_file_type(store, client, name, content):
    r = preview(client, name, content)
    assert r.status_code == 400 and "Unsupported file type" in r.json()["detail"]


def test_fake_xlsx_rejected(store, client):
    r = preview(client, "terms.xlsx", b"term,definition\nx,y")
    assert r.status_code == 400 and "not a valid .xlsx" in r.json()["detail"]


def test_file_size_limit(store, client, monkeypatch):
    monkeypatch.setattr(ti, "MAX_FILE_BYTES", 100)
    r = preview(client, "t.csv", csv_bytes([STRIP_ROW]))
    assert r.status_code == 400 and "too large" in r.json()["detail"]


def test_row_limit(store, client, monkeypatch):
    monkeypatch.setattr(ti, "MAX_ROWS", 3)
    rows = [{**GFE_ROW, "term": f"T{i}", "aliases": ""} for i in range(4)]
    r = preview(client, "t.csv", csv_bytes(rows))
    assert r.status_code == 400 and "Too many rows" in r.json()["detail"]


def test_non_utf8_csv_rejected_cleanly(store, client):
    r = preview(client, "t.csv", "term,definition\ncafé,x".encode("latin-1"))
    assert r.status_code == 400 and "UTF-8" in r.json()["detail"]


def test_xlsx_formulas_not_evaluated(store, client):
    content = xlsx_bytes([["Formula term", "", "", "=1+1", "", "", "", "", "", ""]])
    row = preview(client, "t.xlsx", content).json()["rows"][0]
    # data_only mode: an uncalculated formula has no cached value → definition is blank, never "2"
    assert row["status"] == "invalid" and "Missing required field: definition" in row["errors"]


def test_formula_like_text_in_csv_is_plain_text_with_warning(store, client):
    row = preview(client, "t.csv", csv_bytes([{**GFE_ROW, "details": "=HYPERLINK(\"x\")"}])).json()["rows"][0]
    assert row["status"] == "warning" and row["data"]["details"].startswith("=")


def test_import_requires_owner_key_when_configured(store, client, monkeypatch):
    from rupsaa.config import get_settings

    get_settings.cache_clear()
    monkeypatch.setenv("OWNER_API_KEY", "secret-k")
    get_settings.cache_clear()
    try:
        content = csv_bytes([GFE_ROW])
        assert preview(client, "t.csv", content).status_code == 401
        assert commit(client, "t.csv", content).status_code == 401
        assert store.list() == []
        ok = commit(client, "t.csv", content, headers={"X-Owner-Key": "secret-k"})
        assert ok.status_code == 200 and ok.json()["counts"]["created"] == 1
    finally:
        get_settings.cache_clear()


def test_bad_update_rows_value(store, client):
    r = commit(client, "t.csv", csv_bytes([GFE_ROW]), update_rows="2,abc")
    assert r.status_code == 400


def test_no_filesystem_paths_in_responses(store, client):
    body = commit(client, "C:\\Users\\me\\Desktop\\terms.csv", csv_bytes([GFE_ROW])).text
    assert "Users" not in body and str(store.directory) not in body


# --- imported terms behave like manual ones -------------------------------------------------------

def test_imported_aliases_route_immediately_banglish_bengali_english(store, client):
    xlsx = xlsx_bytes([[STRIP_ROW[c] for c in HEADER], [GFE_ROW[c] for c in HEADER]])
    assert commit(client, "t.xlsx", xlsx).json()["counts"]["created"] == 2
    for msg, expected in [("Strip mane ki?", "term-strip_stripping"),       # Banglish alias
                          ("স্ট্রিপ মানে কী?", "term-strip_stripping"),       # Bengali alias
                          ("what does stripping mean?", "term-strip_stripping"),
                          ("GFE mane ki?", "term-gfe"),
                          ("জিএফই মানে কী?", "term-gfe")]:
        k = build_turn_knowledge(msg, use_rag=False, rag_query=None, terminology=store)
        assert k.route == "terminology" and k.terms_used == [expected], msg


def test_imported_term_editable_deletable_and_disabled_import_respected(store, client):
    commit(client, "t.csv", csv_bytes([STRIP_ROW, {**GFE_ROW, "enabled": "false"}]))
    assert client.put("/owner/terminology/term-strip_stripping", json={"details": "edited"}).json()["details"] == "edited"
    assert store.lookup("GFE mane ki?", classify_message("GFE mane ki?").term_candidate) == []  # disabled
    assert client.delete("/owner/terminology/term-gfe", params={"confirm": True}).status_code == 200
    assert [t.id for t in store.list()] == ["term-strip_stripping"]


def test_existing_single_term_crud_still_works(store, client):
    r = client.post("/owner/terminology", json={"term": "Tip", "definition": "A one-off payment.", "category": "creator_platform"})
    assert r.status_code == 200
    assert client.get("/owner/terminology", params={"q": "tip"}).json()["terms"][0]["id"] == "term-tip"
    assert client.get("/owner/terminology/template.csv").status_code == 200  # not captured by /{term_id}
    assert client.get("/owner/terminology/term-tip").json()["term"] == "Tip"
